"""Semantic checks for the explicitly authored native PPTX retention fixture.

This is a bounded parser for that fixture's converter contract, not a general
Markdown parser or an evaluator of presentation quality or coding agents.
"""

from __future__ import annotations

import copy
import hashlib
import io
import re
import zipfile
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

MAX_MARKDOWN_BYTES = 1024 * 1024
# Limits apply to the fixed three-slide control, not arbitrary presentations.
MAX_PACKAGE_BYTES = 1024 * 1024
MAX_MEMBER_BYTES = 1024 * 1024
MAX_EXPANDED_BYTES = 4 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 512
MAX_PACKAGE_DEPTH = 1
MAX_CORE_XML_BYTES = 64 * 1024
MAX_CORE_XML_NODES = 4096
MAX_CORE_XML_DEPTH = 64
SLIDE_HEADING = re.compile(r"^## Slide\s+(\d+)\s*$", re.MULTILINE)


@dataclass
class _ArchiveBudget:
    """Accumulate limits across the outer presentation and embedded workbooks."""

    members: int = 0
    expanded_bytes: int = 0


def _check_member_name(info: zipfile.ZipInfo) -> None:
    """Reject names that could escape an extraction root in downstream tools."""
    name = info.filename
    parts = name.removesuffix("/").split("/")
    if (
        not name
        or info.orig_filename != name
        or name.startswith("/")
        or "\\" in name
        or any(part in {"", ".", ".."} for part in parts)
        or ":" in parts[0]
    ):
        raise ValueError("SDK fixture contains an unsafe ZIP member name")


def _normalize_core_properties(data: bytes) -> bytes:
    """Normalize bounded core properties without DTDs, entities or excessive structure."""
    from defusedxml import ElementTree

    if len(data) > MAX_CORE_XML_BYTES:
        raise ValueError("SDK fixture core properties exceed the XML byte budget")
    root = ElementTree.fromstring(data, forbid_dtd=True, forbid_entities=True, forbid_external=True)
    pending = [(root, 1)]
    visited = 0
    while pending:
        node, depth = pending.pop()
        visited += 1
        if visited > MAX_CORE_XML_NODES or depth > MAX_CORE_XML_DEPTH:
            raise ValueError("SDK fixture core properties exceed the XML structure budget")
        pending.extend((child, depth + 1) for child in node)
    for key in ("created", "modified"):
        for field in root.findall("{http://purl.org/dc/terms/}" + key):
            field.text = "2026-01-01T00:00:00Z"
    return ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)


def _read_member(source: zipfile.ZipFile, member: zipfile.ZipInfo) -> bytes:
    """Admit a supported ZIP member and read at most the per-member budget."""
    _check_member_name(member)
    if member.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}:
        raise ValueError("SDK fixture uses an unsupported ZIP compression method")
    if member.flag_bits & 1:
        raise ValueError("SDK fixture contains an encrypted ZIP member")
    if member.file_size > MAX_MEMBER_BYTES:
        raise ValueError("SDK fixture member exceeds the byte budget")
    with source.open(member) as stream:
        data = stream.read(MAX_MEMBER_BYTES + 1)
    if len(data) != member.file_size:
        raise ValueError("SDK fixture member size does not match the ZIP directory")
    return data


def _tables(text: str) -> list[list[list[str]]]:
    tables = []
    current: list[list[str]] = []
    for line in [*text.splitlines(), ""]:
        line = line.strip()
        if line.startswith("|") and line.endswith("|"):
            cells = [cell.strip().replace(r"\|", "|") for cell in re.split(r"(?<!\\)\|", line[1:-1])]
            if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                current.append(cells)
        elif current:
            tables.append(current)
            current = []
    return tables


def _table(text: str, header: list[str]) -> list[list[str]] | None:
    matching = [rows[1:] for rows in _tables(text) if rows and rows[0] == header]
    return matching[0] if len(matching) == 1 else None


def _number(value: str, expected: str) -> bool:
    try:
        actual = Decimal(value)
    except InvalidOperation:
        return False
    return actual.is_finite() and actual == Decimal(expected)


def _category_values(text: str) -> bool:
    rows = _table(text, ["Category", "Bar_series"])
    if rows is None or len(rows) != 2 or any(len(row) != 2 for row in rows):
        return False
    return (
        rows[0][0] == "North_Unit"
        and _number(rows[0][1], "0")
        and rows[1][0] == "South_Unit"
        and _number(rows[1][1], "-7.125")
    )


def _xy_values(text: str) -> bool:
    rows = _table(text, ["Series", "Point", "X", "Y"])
    if rows is None or len(rows) != 2 or any(len(row) != 4 for row in rows):
        return False
    expected = [("1", "1.25", "-2.5"), ("2", "2.75", "0")]
    return all(
        row[0] == "XY_series_control"
        and all(_number(value, wanted) for value, wanted in zip(row[1:], coordinates, strict=True))
        for row, coordinates in zip(rows, expected, strict=True)
    )


def evaluate_markdown(text: str) -> dict[str, bool]:
    """Compare slide-local content and numeric associations with authored facts."""
    if len(text.encode("utf-8")) > MAX_MARKDOWN_BYTES:
        raise ValueError("retention Markdown exceeds the one MiB evaluation budget")
    headings = list(SLIDE_HEADING.finditer(text))
    ids = [int(match.group(1)) for match in headings]
    slides = {}
    for index, match in enumerate(headings):
        end = headings[index + 1].start() if index + 1 < len(headings) else len(text)
        slides[int(match.group(1))] = text[match.end() : end]
    first, second, third = (slides.get(index, "") for index in (1, 2, 3))
    body, marker, notes = first.partition("### Speaker Notes")
    return {
        "unicode_text": "Unicode control Alpha β Omega Ω" in body,
        "table_cells": _table(body, ["Table_header_A", "Table_header_B"]) == [["Table_body_C", "Table_body_D"]],
        "speaker_notes": bool(marker) and "Speaker_note_control_K26" in notes,
        "grouped_text": "Grouped_text_control_δ" in second,
        "hyperlink": "[Reference_link_control](https://example.invalid/evidence)" in body,
        "category_numeric_value": _category_values(second),
        "xy_coordinates": _xy_values(second),
        "office14_choice_text": "Office14_choice_text_control" in third,
        "slide_heading_count": ids == [1, 2, 3],
    }


def normalize_package(raw: bytes) -> bytes:
    """Freeze outer and embedded workbook ZIP dates and core property dates."""
    return _normalize_package(raw, _ArchiveBudget(), 0)


def _normalize_package(raw: bytes, budget: _ArchiveBudget, depth: int) -> bytes:
    """Apply shared archive limits and normalize a single supported ZIP layer."""
    if depth > MAX_PACKAGE_DEPTH:
        raise ValueError("SDK fixture contains nested workbooks beyond the supported depth")
    if len(raw) > MAX_PACKAGE_BYTES:
        raise ValueError("SDK fixture package exceeds the compressed byte budget")
    result = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(raw)) as source,
        zipfile.ZipFile(result, "w", compression=zipfile.ZIP_DEFLATED) as target,
    ):
        members = source.infolist()
        names = [member.filename for member in members]
        if len(names) != len(set(names)):
            raise ValueError("SDK fixture contains duplicate ZIP members")
        budget.members += len(members)
        budget.expanded_bytes += sum(member.file_size for member in members)
        if budget.members > MAX_ARCHIVE_MEMBERS or budget.expanded_bytes > MAX_EXPANDED_BYTES:
            raise ValueError("SDK fixture exceeds the archive member or expanded byte budget")
        for member in sorted(members, key=lambda item: item.filename):
            data = _read_member(source, member)
            name = member.filename
            if name.lower().endswith(".xlsx"):
                data = _normalize_package(data, budget, depth + 1)
            elif name == "docProps/core.xml":
                data = _normalize_core_properties(data)
            info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o100644 << 16
            target.writestr(info, data)
    return result.getvalue()


def create_fixture(directory: Path) -> tuple[Path, dict[str, Any]]:
    """Create the SDK-authored fixture; optional consumer packages load here only."""
    from lxml import etree
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData, XyChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    directory.mkdir(parents=True, exist_ok=False)
    fixture = directory / "retention.pptx"
    raw_fixture = directory / "retention-sdk-output.pptx"
    presentation = Presentation()
    presentation.core_properties.created = datetime(2026, 1, 1)
    presentation.core_properties.modified = datetime(2026, 1, 1)
    presentation.core_properties.author = "Benchmark fixture"
    presentation.core_properties.last_modified_by = "Benchmark fixture"
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "Unicode control Alpha β Omega Ω"
    slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1)).text = "## Slide 33"
    link = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(6), Inches(1)).text_frame.paragraphs[0].add_run()
    link.text = "Reference_link_control"
    link.hyperlink.address = "https://example.invalid/evidence"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(6), Inches(1)).table
    cells = ["Table_header_A", "Table_header_B", "Table_body_C", "Table_body_D"]
    for index, value in enumerate(cells):
        table.cell(index // 2, index % 2).text = value
    slide.notes_slide.notes_text_frame.text = "Speaker_note_control_K26"
    chart_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    bar = CategoryChartData()
    bar.categories = ["North_Unit", "South_Unit"]
    bar.add_series("Bar_series", [0, -7.125])
    bar_shape = chart_slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(0.5), Inches(4), Inches(3), bar
    )
    bar_shape.name = "Bar_retention_control"
    xy = XyChartData()
    series = xy.add_series("XY_series_control")
    series.add_data_point(1.25, -2.5)
    series.add_data_point(2.75, 0)
    xy_shape = chart_slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(5), Inches(0.5), Inches(4), Inches(3), xy)
    xy_shape.name = "XY_retention_control"
    group = chart_slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(4), Inches(5), Inches(1)).text = "Grouped_text_control_δ"
    choice_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    preferred = choice_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    preferred.text = "Office14_choice_text_control"
    element = preferred.element
    parent = element.getparent()
    position = parent.index(element)
    parent.remove(element)
    mc = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    a14 = "http://schemas.microsoft.com/office/drawing/2010/main"
    alternate = etree.Element("{" + mc + "}AlternateContent", nsmap={"mc": mc, "a14": a14})
    choice = etree.SubElement(alternate, "{" + mc + "}Choice", Requires="a14")
    choice.append(element)
    fallback = etree.SubElement(alternate, "{" + mc + "}Fallback")
    fallback_shape = copy.deepcopy(element)
    ns = {
        "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    }
    fallback_shape.find("p:nvSpPr/p:cNvPr", ns).set("id", "9000")
    for text in fallback_shape.findall(".//a:t", ns):
        text.text = "Fallback_text_control"
    fallback.append(fallback_shape)
    parent.insert(position, alternate)
    presentation.save(raw_fixture)
    fixture.write_bytes(normalize_package(raw_fixture.read_bytes()))
    task = {
        "schema_version": 1,
        "evidence_class": "native-component-content-retention-controls",
        "task_definition": "Convert the same SDK-authored three-slide fixture to Markdown while retaining labeled text, table cells, hyperlinks, chart values, notes and the supported Office14 Choice text; body text must not create extra slide headings.",
        "fixture_sha256": hashlib.sha256(fixture.read_bytes()).hexdigest(),
        "expected_slide_count": 3,
        "expected_table_cells": cells,
        "expected_xy_x_values": ["1.25", "2.75"],
        "expected_xy_y_values": ["-2.5", "0"],
        "expected_category_values": ["0", "-7.125"],
        "definition_independent_of_exact_skill_text": True,
        "expected_results_basis": "The explicitly authored input labels and values, plus the number of physical slide parts in its PPTX ZIP. No expected answer was extracted from SKILL.md.",
        "coding_agent_efficacy_scored": False,
        "source_workflow_completion_scored": False,
    }
    return fixture, task
